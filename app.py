import streamlit as st
import os
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from fpdf import FPDF
from pdf2docx import Converter
import qrcode
import openai

# Set OpenAI API Key
openai.api_key = st.secrets["OPENAI_API_KEY"]

# Function: DOCX to PDF
def docx_to_pdf(input_path, output_path):
    document = Document(input_path)
    pdf = canvas.Canvas(output_path, pagesize=letter)
    y_position = 750  
    for para in document.paragraphs:
        pdf.drawString(100, y_position, para.text)
        y_position -= 20  
        if y_position < 50:
            pdf.showPage()
            y_position = 750
    pdf.save()

# Function: PDF to DOCX
def pdf_to_docx(input_path, output_path):
    cv = Converter(input_path)
    cv.convert(output_path, start=0, end=None)
    cv.close()

# Function: Text to PDF
def text_to_pdf(text, output_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    for line in text.split('\n'):
        pdf.cell(200, 10, txt=line, ln=True)
    pdf.output(output_path)

# Function: Image to QR Code
def image_to_qr_code(image_url, output_path):
    qr = qrcode.make(image_url)
    qr.save(output_path)

# Function: Generate AI Image
def generate_ai_image(prompt):
    try:
        response = openai.Image.create(
            prompt=prompt,
            n=1,
            size="512x512"
        )
        image_url = response['data'][0]['url']
        return image_url
    except Exception as e:
        return f"Error generating image: {e}"

# Streamlit UI
def main():
    st.title("🔥 QuickToolBox - All-in-One Tool 🔥")

    st.sidebar.title("Choose a Tool")
    option = st.sidebar.selectbox(
        "Select an action",
        (
            "DOCX to PDF",
            "PDF to DOCX",
            "Text to PDF",
            "Image URL to QR Code",
            "AI Image Generator"
        )
    )

    if option == "DOCX to PDF":
        st.header("📄 DOCX to PDF Converter")
        uploaded_file = st.file_uploader("Upload a DOCX file", type="docx")
        if uploaded_file:
            with open("uploaded.docx", "wb") as f:
                f.write(uploaded_file.getbuffer())
            output_path = "converted.pdf"
            docx_to_pdf("uploaded.docx", output_path)
            st.success("Conversion successful!")
            with open(output_path, "rb") as file:
                st.download_button(
                    label="Download PDF",
                    data=file,
                    file_name="converted.pdf",
                    mime="application/pdf"
                )
            os.remove("uploaded.docx")

    elif option == "PDF to DOCX":
        st.header("📄 PDF to DOCX Converter")
        uploaded_file = st.file_uploader("Upload a PDF file", type="pdf")
        if uploaded_file:
            with open("uploaded.pdf", "wb") as f:
                f.write(uploaded_file.getbuffer())
            output_path = "converted.docx"
            pdf_to_docx("uploaded.pdf", output_path)
            st.success("Conversion successful!")
            with open(output_path, "rb") as file:
                st.download_button(
                    label="Download DOCX",
                    data=file,
                    file_name="converted.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
            os.remove("uploaded.pdf")

    elif option == "Text to PDF":
        st.header("📝 Text to PDF Converter")
        text_input = st.text_area("Enter your text")
        if st.button("Convert to PDF"):
            output_path = "text_output.pdf"
            text_to_pdf(text_input, output_path)
            st.success("Conversion successful!")
            with open(output_path, "rb") as file:
                st.download_button(
                    label="Download PDF",
                    data=file,
                    file_name="text_output.pdf",
                    mime="application/pdf"
                )

    elif option == "Image URL to QR Code":
        st.header("🌐 Image URL to QR Code Generator")
        image_url = st.text_input("Enter Image URL")
        if st.button("Generate QR Code"):
            output_path = "qrcode.png"
            image_to_qr_code(image_url, output_path)
            st.image(output_path, caption="Generated QR Code", use_column_width=True)
            with open(output_path, "rb") as file:
                st.download_button(
                    label="Download QR Code",
                    data=file,
                    file_name="qrcode.png",
                    mime="image/png"
                )

    elif option == "AI Image Generator":
        st.header("🎨 AI Image Generator")
        prompt = st.text_input("Enter a prompt to generate an image")
        if st.button("Generate Image"):
            img_url = generate_ai_image(prompt)
            if "Error" not in img_url:
                st.image(img_url, caption="Generated Image", use_column_width=True)
            else:
                st.error(img_url)

if __name__ == "__main__":
    main()
